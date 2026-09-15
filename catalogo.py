from openpyxl import load_workbook
from urllib.parse import quote
from pathlib import Path
import unicodedata
import re
import os
import requests

SUPABASE_URL = os.environ.get("SUPABASE_URL", "").rstrip("/")
SUPABASE_ANON_KEY = os.environ.get("SUPABASE_ANON_KEY", "")

WHATSAPP_NUMERO = "5491163509142"  # Celular WhatsApp sin + ni espacios.
ARCHIVO_PRODUCTOS = "productos.xlsx"

MARCAS_ORDEN = ["Ana Grant", "Aretha", "BLUO", "Deville", "So Pink", "Stylo"]

GRUPOS_CATEGORIAS = {
    "Lencería": ["Ropa de dormir", "Medias"],
    "Corsetería": [
        "Corpiños",
        "Bombachas",
        "Portaligas y ligas",
        "Corset",
        "Body",
        "Bra importado",
    ],
}


def normalizar(texto):
    texto = str(texto or "").strip()
    texto = unicodedata.normalize("NFKD", texto)
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    return texto.lower()


def contiene(texto, palabras):
    texto_normalizado = normalizar(texto)
    return any(palabra in texto_normalizado for palabra in palabras)


def marca_carpeta(marca):
    mapa = {
        "ana grant": "Ana Grant",
        "aretha": "Aretha",
        "bluo": "BLUO",
        "deville": "Deville",
        "so pink": "So Pink",
        "so pink!": "So Pink",
        "stylo": "Stylo",
    }
    return mapa.get(normalizar(marca), str(marca).strip())


def es_destacado(valor):
    return normalizar(valor) in ["si", "sí", "s", "x", "true", "1", "destacado"]


def obtener_descuento(valor):
    """
    Lee la columna Oferta del Excel.
    Ejemplos válidos: "oferta 10%", "oferta 15", "20%" o "20".
    """
    texto = normalizar(valor)
    if not texto:
        return 0

    match = re.search(r"(\d{1,2})", texto)
    if not match:
        return 0

    porcentaje = int(match.group(1))
    if porcentaje <= 0 or porcentaje >= 100:
        return 0

    return porcentaje


def precio_a_numero(precio):
    texto = str(precio or "").strip()
    if not texto:
        return None

    texto = texto.replace("$", "").replace(" ", "")

    if "," in texto:
        texto = texto.replace(".", "").replace(",", ".")
    else:
        texto = texto.replace(".", "")

    try:
        return float(texto)
    except ValueError:
        return None


def formatear_pesos(valor):
    if valor is None:
        return ""

    entero = int(round(valor))
    return "$" + f"{entero:,}".replace(",", ".")


def calcular_precio_oferta(precio, descuento):
    numero = precio_a_numero(precio)
    if numero is None or descuento <= 0:
        return ""

    return formatear_pesos(numero * (1 - descuento / 100))


def precio_html(producto):
    descuento = producto.get("descuento", 0)
    precio = producto.get("precio", "")

    if descuento > 0:
        precio_oferta = producto.get("precio_oferta") or calcular_precio_oferta(precio, descuento)
        return f"""
          <p class="precio precio-con-oferta">
            <span class="precio-anterior">{precio}</span>
            <span class="precio-oferta">{precio_oferta}</span>
            <span class="descuento-pill">−{descuento}%</span>
          </p>
        """

    return f'<p class="precio">{precio}</p>'


def producto_visible_por_stock(valor):
    """
    Controla desde Excel si un producto se muestra o no.

    En la columna Stock podés poner:
    - "No", "Sin stock", "Agotado", "No disponible" o "0" para ocultarlo.
    - "Sí", "Disponible", un número mayor a 0 o dejarlo vacío para mostrarlo.

    También oculta variantes como "SIN STOCK", "sin-stock",
    "agotado momentáneamente" o "no disponible".
    """
    if valor is None:
        return True

    if isinstance(valor, (int, float)):
        return valor > 0

    stock = normalizar(valor)
    stock_simple = re.sub(r"[^a-z0-9]+", " ", stock).strip()

    if stock_simple == "":
        return True

    ocultar_exactos = {
        "0",
        "no",
        "n",
        "false",
        "falso",
        "sin stock",
        "stock 0",
        "no stock",
        "no disponible",
        "no hay stock",
    }

    ocultar_si_contiene = [
        "sin stock",
        "agotado",
        "agotada",
        "agotados",
        "agotadas",
        "no disponible",
        "discontinuado",
        "discontinuada",
    ]

    if stock_simple in ocultar_exactos:
        return False

    if any(texto in stock_simple for texto in ocultar_si_contiene):
        return False

    return True


def clasificar_producto(producto):
    """
    Clasifica cada producto usando las columnas B, C y D del Excel:
    Nombre, Descripción y Rubro.
    """
    texto = " ".join([
        producto.get("nombre", ""),
        producto.get("descripcion", ""),
        producto.get("rubro", ""),
    ])

    reglas = [
        ("Lencería", "Ropa de dormir", [
            "camison", "bata", "pijama", "pijamas", "pantufla", "pantuflas",
            "pantuflon", "remeron", "musculosa", "remera", "short",
            "camiseta", "bermuda"
        ]),
        ("Lencería", "Medias", [
            "casual", "soquete", "canoa", "media", "medias", "invisible",
            "manguita"
        ]),
        ("Corsetería", "Portaligas y ligas", [
            "portaliga", "portaligas", "liga", "ligas"
        ]),
        ("Corsetería", "Corset", [
            "corset", "faja"
        ]),
        ("Corsetería", "Body", [
            "body", "bodies"
        ]),
        ("Corsetería", "Bra importado", [
            "tasa de silicona", "tasas de silicona", "taza de silicona",
            "tazas de silicona", "tasa de siliconas", "tasas de siliconas",
            "taza de siliconas", "tazas de siliconas", "pesonera",
            "pesoneras", "boop tape", "boob tape", "body tape"
        ]),
        ("Corsetería", "Bombachas", [
            "bikini", "colaless", "cola les", "culotte", "culote",
            "culotteless", "bombacha", "bombachas", "trusa", "brief",
            "vedetina", "tiro corto", "tanga"
        ]),
        ("Corsetería", "Corpiños", [
            "soutien", "soutiens", "corpiño", "corpino", "corpiños",
            "corpinos", "taza", "tazas", "triangulo", "triángulo",
            "bandeau"
        ]),
    ]

    for grupo, categoria, palabras in reglas:
        if contiene(texto, palabras):
            return grupo, categoria

    return "Otros", "Otros"


def buscar_imagen(codigo, marca):
    carpeta = Path("imagenes") / marca_carpeta(marca)
    if not carpeta.exists():
        return "imagenes/logo-mucho-de-ti.jpg"

    codigo_limpio = str(codigo).strip().lower()
    extensiones = ["*.jpg", "*.jpeg", "*.png", "*.webp"]

    for patron in extensiones:
        for archivo in carpeta.glob(patron):
            nombre = archivo.name.lower().strip()
            # Vincula archivos como: "109. Ana Grant.jpg", "109 Ana Grant.jpg", "109-Ana Grant.jpg"
            if re.match(rf"^{re.escape(codigo_limpio)}(\D|$)", nombre):
                return archivo.as_posix()

    return "imagenes/logo-mucho-de-ti.jpg"


def leer_productos_excel():
    wb = load_workbook(ARCHIVO_PRODUCTOS, data_only=True)
    ws = wb.active
    encabezados = [str(c.value or "").strip().lower() for c in ws[1]]

    def col(nombre, requerido=True):
        nombre = nombre.lower()
        if nombre in encabezados:
            return encabezados.index(nombre)
        if requerido:
            raise ValueError(f"Falta la columna obligatoria '{nombre}' en {ARCHIVO_PRODUCTOS}")
        return None

    productos = []
    for fila in ws.iter_rows(min_row=2, values_only=True):
        if not any(fila):
            continue

        producto = {
            "codigo": str(fila[col("codigo")] or "").strip(),
            "nombre": str(fila[col("nombre")] or "").strip(),
            "descripcion": str(fila[col("descripcion")] or "").strip(),
            "rubro": str(fila[col("rubro")] or "").strip(),
            "marca": marca_carpeta(fila[col("marca")]),
            "precio": str(fila[col("precio")] or "").strip(),
            "talles": str(fila[col("talles")] or "Consultar disponibilidad").strip(),
            "stock": str(fila[col("stock")] or "").strip(),
            "destacado": es_destacado(fila[col("destacado")]),
        }

        oferta_col = col("oferta", requerido=False)
        oferta_valor = fila[oferta_col] if oferta_col is not None else ""
        descuento = obtener_descuento(oferta_valor)

        producto["oferta"] = str(oferta_valor or "").strip()
        producto["descuento"] = descuento
        producto["precio_oferta"] = calcular_precio_oferta(producto["precio"], descuento)

        if producto["codigo"] and producto["nombre"]:
            grupo, categoria = clasificar_producto(producto)
            producto["grupo"] = grupo
            producto["categoria"] = categoria
            producto["imagen"] = buscar_imagen(producto["codigo"], producto["marca"])
            productos.append(producto)

    return productos


def url_imagen_supabase(imagen_path):
    if not imagen_path:
        return "imagenes/logo-mucho-de-ti.jpg"
    return f"{SUPABASE_URL}/storage/v1/object/public/imagenes/{quote(str(imagen_path))}"


def leer_productos_supabase():
    """
    Fuente de datos real del sitio: la tabla `productos` de Supabase,
    editada desde el panel de administración (ver supabase/schema.sql).
    productos.xlsx queda como respaldo histórico.
    """
    respuesta = requests.get(
        f"{SUPABASE_URL}/rest/v1/productos",
        params={"select": "*"},
        headers={
            "apikey": SUPABASE_ANON_KEY,
            "Authorization": f"Bearer {SUPABASE_ANON_KEY}",
        },
        timeout=30,
    )
    respuesta.raise_for_status()

    productos = []
    for fila in respuesta.json():
        producto = {
            "codigo": str(fila.get("codigo") or "").strip(),
            "nombre": str(fila.get("nombre") or "").strip(),
            "descripcion": str(fila.get("descripcion") or "").strip(),
            "rubro": str(fila.get("rubro") or "").strip(),
            "marca": marca_carpeta(fila.get("marca")),
            "precio": str(fila.get("precio") or "").strip(),
            "talles": str(fila.get("talles") or "Consultar disponibilidad").strip(),
            "stock": str(fila.get("stock") or "").strip(),
            "destacado": bool(fila.get("destacado")),
        }

        oferta_valor = fila.get("oferta")
        descuento = obtener_descuento(oferta_valor)
        producto["oferta"] = str(oferta_valor or "").strip()
        producto["descuento"] = descuento
        producto["precio_oferta"] = calcular_precio_oferta(producto["precio"], descuento)

        if producto["codigo"] and producto["nombre"]:
            grupo, categoria = clasificar_producto(producto)
            producto["grupo"] = grupo
            producto["categoria"] = categoria
            producto["imagen"] = url_imagen_supabase(fila.get("imagen_path"))
            productos.append(producto)

    return productos
